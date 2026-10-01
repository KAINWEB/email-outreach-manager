import os, sys, sqlite3, csv, io, re, json, smtplib, ssl, threading, time, webbrowser
from datetime import datetime
from email.message import EmailMessage
from fastapi import FastAPI, UploadFile, File, HTTPException
from fastapi.staticfiles import StaticFiles
from fastapi.responses import FileResponse, StreamingResponse
from pydantic import BaseModel
from openpyxl import load_workbook
import keyring

BASE_DIR=os.path.dirname(sys.executable) if getattr(sys,"frozen",False) else os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA_DIR=os.path.join(BASE_DIR,"data"); os.makedirs(DATA_DIR,exist_ok=True)
DB=os.path.join(DATA_DIR,"app.db"); SERVICE="GrowerTeam.EmailOutreachManager"
UI=os.path.join(getattr(sys,"_MEIPASS",BASE_DIR),"frontend_dist")
app=FastAPI(title="GROWER TEAM Email Outreach Manager")
pause_event=threading.Event(); pause_event.set(); worker_stop=threading.Event()

def con():
    c=sqlite3.connect(DB,check_same_thread=False); c.row_factory=sqlite3.Row; return c
def init():
    with con() as c:
        c.executescript("""CREATE TABLE IF NOT EXISTS senders(id INTEGER PRIMARY KEY,email TEXT UNIQUE,name TEXT,host TEXT,port INTEGER,use_ssl INTEGER DEFAULT 0,use_starttls INTEGER DEFAULT 1,enabled INTEGER DEFAULT 1,daily_limit INTEGER DEFAULT 100,rate_seconds REAL DEFAULT 3,sent_today INTEGER DEFAULT 0,last_error TEXT);
CREATE TABLE IF NOT EXISTS campaigns(id INTEGER PRIMARY KEY,name TEXT,status TEXT DEFAULT 'DRAFT',created_at TEXT);
CREATE TABLE IF NOT EXISTS recipients(id INTEGER PRIMARY KEY,campaign_id INTEGER,email TEXT,company TEXT,name TEXT,subject TEXT,message TEXT,status TEXT DEFAULT 'READY',sender_id INTEGER,error TEXT,attempts INTEGER DEFAULT 0,created_at TEXT);
CREATE TABLE IF NOT EXISTS suppression(email TEXT PRIMARY KEY,reason TEXT,created_at TEXT);
CREATE TABLE IF NOT EXISTS logs(id INTEGER PRIMARY KEY,ts TEXT,campaign_id INTEGER,recipient_id INTEGER,sender TEXT,email TEXT,subject TEXT,status TEXT,error TEXT);""")
def rows(q,p=()):
    with con() as c:return [dict(x) for x in c.execute(q,p).fetchall()]
def one(q,p=()):
    with con() as c:
        x=c.execute(q,p).fetchone(); return dict(x) if x else None
def valid_email(s): return bool(re.match(r"^[^@\s]+@[^@\s]+\.[^@\s]+$",s or ""))
def log(cid,rid,sender,email,subject,status,error=""):
    with con() as c:c.execute("INSERT INTO logs(ts,campaign_id,recipient_id,sender,email,subject,status,error) VALUES(?,?,?,?,?,?,?,?)",(datetime.now().isoformat(timespec="seconds"),cid,rid,sender,email,subject,status,error))
class SenderIn(BaseModel):
    email:str; name:str=""; host:str; port:int=587; use_ssl:bool=False; use_starttls:bool=True; password:str=""; daily_limit:int=100; rate_seconds:float=3
class CampaignIn(BaseModel): name:str
class MapIn(BaseModel): campaign_id:int; mapping:dict
class IdIn(BaseModel): id:int
class SuppIn(BaseModel): email:str; reason:str="manual"
class RecipientEdit(BaseModel): email:str; company:str=""; name:str=""; subject:str; message:str
@app.get("/api/health")
def health(): return {"status":"ok","version":"2.0.0"}
@app.get("/api/dashboard")
def dashboard():
    return {"senders":one("SELECT COUNT(*) n FROM senders")["n"],"campaigns":one("SELECT COUNT(*) n FROM campaigns")["n"],"ready":one("SELECT COUNT(*) n FROM recipients WHERE status IN ('READY','QUEUED')")["n"],"sent":one("SELECT COUNT(*) n FROM recipients WHERE status='SENT'")["n"],"failed":one("SELECT COUNT(*) n FROM recipients WHERE status='FAILED'")["n"],"paused":not pause_event.is_set()}
@app.get("/api/senders")
def senders(): return rows("SELECT id,email,name,host,port,use_ssl,use_starttls,enabled,daily_limit,rate_seconds,sent_today,last_error FROM senders ORDER BY id DESC")
@app.post("/api/senders")
def add_sender(x:SenderIn):
    with con() as c:
        cur=c.execute("INSERT INTO senders(email,name,host,port,use_ssl,use_starttls,daily_limit,rate_seconds) VALUES(?,?,?,?,?,?,?,?)",(x.email,x.name,x.host,x.port,int(x.use_ssl),int(x.use_starttls),x.daily_limit,x.rate_seconds)); sid=cur.lastrowid
    if x.password:keyring.set_password(SERVICE,x.email,x.password)
    return {"id":sid}
def smtp_connect(s):
    pwd=keyring.get_password(SERVICE,s["email"])
    if not pwd: raise RuntimeError("Пароль приложения не сохранён")
    if s["use_ssl"]: server=smtplib.SMTP_SSL(s["host"],s["port"],timeout=20,context=ssl.create_default_context())
    else:
        server=smtplib.SMTP(s["host"],s["port"],timeout=20)
        if s["use_starttls"]: server.starttls(context=ssl.create_default_context())
    server.login(s["email"],pwd); return server
@app.post("/api/senders/{sid}/test")
def test_sender(sid:int):
    s=one("SELECT * FROM senders WHERE id=?",(sid,))
    if not s: raise HTTPException(404)
    try:
        z=smtp_connect(s); z.quit(); err=""
    except Exception as e: err=str(e)
    with con() as c:c.execute("UPDATE senders SET last_error=? WHERE id=?",(err,sid))
    return {"ok":not err,"error":err}
@app.post("/api/senders/{sid}/toggle")
def toggle_sender(sid:int):
    with con() as c:c.execute("UPDATE senders SET enabled=1-enabled WHERE id=?",(sid,))
    return {"ok":True}
@app.delete("/api/senders/{sid}")
def del_sender(sid:int):
    s=one("SELECT email FROM senders WHERE id=?",(sid,))
    with con() as c:c.execute("DELETE FROM senders WHERE id=?",(sid,))
    if s:
        try:keyring.delete_password(SERVICE,s["email"])
        except:pass
    return {"ok":True}
@app.get("/api/campaigns")
def campaigns(): return rows("SELECT c.*,COUNT(r.id) total,SUM(CASE WHEN r.status='SENT' THEN 1 ELSE 0 END) sent,SUM(CASE WHEN r.status='FAILED' THEN 1 ELSE 0 END) failed FROM campaigns c LEFT JOIN recipients r ON r.campaign_id=c.id GROUP BY c.id ORDER BY c.id DESC")
@app.post("/api/campaigns")
def campaign(x:CampaignIn):
    with con() as c: cur=c.execute("INSERT INTO campaigns(name,created_at) VALUES(?,?)",(x.name,datetime.now().isoformat(timespec="seconds"))); return {"id":cur.lastrowid}
def parse_upload(data,filename):
    ext=(filename or "").lower()
    if ext.endswith(".xlsx"):
        wb=load_workbook(io.BytesIO(data),read_only=True,data_only=True); ws=wb.active
        vals=list(ws.iter_rows(values_only=True))
        if not vals:return [],[]
        hs=[str(x or "").strip() for x in vals[0]]
        return hs,[dict(zip(hs,r)) for r in vals[1:]]
    text=data.decode("utf-8-sig",errors="replace")
    try:dialect=csv.Sniffer().sniff(text[:8192],delimiters=",;\t")
    except: dialect=csv.excel
    rr=list(csv.DictReader(io.StringIO(text),dialect=dialect))
    return (list(rr[0].keys()) if rr else []),rr

def norm(s): return re.sub(r"[^a-zа-я0-9]+"," ",str(s or "").lower()).strip()
def detect_columns(headers):
    nh={h:norm(h) for h in headers}
    def pick(exacts=(),contains=()):
        for h,n in nh.items():
            if n in exacts:return h
        for h,n in nh.items():
            if any(x in n for x in contains):return h
        return ""
    return {
      "email":pick(("email","e mail","рабочий email контакт","личный email контакт"),("рабочий email","email контакт","электронная почта","e mail")),
      "company":pick(("company","компания","название компания"),("название компания","company")),
      "name":pick(("name","имя","полное имя контакт"),("полное имя","фио","contact name")),
      "subject":pick(("subject","тема","тема письма"),("тема письма","email subject")),
      "message":pick(("message","сообщение","текст письма","примечание к контакту"),("примечание к контакту","текст письма","сообщение")),
      "deal":pick(("название сделка","deal"),("название сделка","сделк"))
    }
def make_subject(r,mp):
    if mp.get("subject") and str(r.get(mp["subject"],"") or "").strip():return str(r.get(mp["subject"])).strip()
    company=str(r.get(mp.get("company",""),"") or "").strip()
    deal=str(r.get(mp.get("deal",""),"") or "").strip()
    return ("Идея по развитию "+company) if company else (deal or "Предложение по развитию")
def import_rows(cid,source,mp):
    added=skipped=duplicates=invalid=0; seen=set()
    with con() as c:
        suppressed={x[0].lower() for x in c.execute("SELECT email FROM suppression")}
        existing={x[0].lower() for x in c.execute("SELECT email FROM recipients WHERE campaign_id=?",(cid,))}
        for r in source:
            get=lambda k:str(r.get(mp.get(k,""),"") or "").strip()
            email=get("email").lower(); message=get("message"); subject=make_subject(r,mp)
            if email in seen or email in existing: skipped+=1;duplicates+=1;continue
            if not valid_email(email) or not message: skipped+=1;invalid+=1;continue
            if email in suppressed: skipped+=1;continue
            seen.add(email)
            name=get("name"); name="" if norm(name) in ("нету","нет","none","nan") else name
            c.execute("INSERT INTO recipients(campaign_id,email,company,name,subject,message,status,created_at) VALUES(?,?,?,?,?,?,?,?)",(cid,email,get("company"),name,subject,message,"READY",datetime.now().isoformat(timespec="seconds")));added+=1
    return {"added":added,"skipped":skipped,"duplicates":duplicates,"invalid":invalid}

@app.post("/api/import/auto")
async def auto_import(file:UploadFile=File(...)):
    data=await file.read(); hs,source=parse_upload(data,file.filename); mp=detect_columns(hs)
    if not mp["email"] or not mp["message"]: raise HTTPException(400,detail={"message":"Не удалось автоматически найти email или текст письма","detected":mp,"headers":hs})
    base=os.path.splitext(file.filename or "Импорт")[0]
    with con() as c:
        cur=c.execute("INSERT INTO campaigns(name,status,created_at) VALUES(?,?,?)",(base,"DRAFT",datetime.now().isoformat(timespec="seconds")));cid=cur.lastrowid
    result=import_rows(cid,source,mp)
    senders=rows("SELECT id,email,name,daily_limit,sent_today,rate_seconds,enabled FROM senders WHERE enabled=1 ORDER BY id")
    capacity=sum(max(0,int(s["daily_limit"])-int(s["sent_today"])) for s in senders)
    remaining=result["added"]; allocation=[]
    for s in senders:
        free=max(0,int(s["daily_limit"])-int(s["sent_today"])); take=min(free,remaining)
        if take: allocation.append({"sender":s["email"],"count":take,"rate_seconds":s["rate_seconds"]});remaining-=take
    return {"campaign_id":cid,"campaign":base,"detected":mp,**result,"active_senders":len(senders),"today_capacity":capacity,"planned_today":min(result["added"],capacity),"remaining_after_today":max(0,result["added"]-capacity),"allocation":allocation}

@app.post("/api/import/preview")
async def preview(file:UploadFile=File(...)):
    data=await file.read();hs,source=parse_upload(data,file.filename);mp=detect_columns(hs)
    return {"headers":hs,"sample":source[:5],"detected":mp}

@app.post("/api/import/{cid}")
async def do_import(cid:int,file:UploadFile=File(...),mapping:str="{}"):
    data=await file.read();hs,source=parse_upload(data,file.filename); supplied=json.loads(mapping or "{}")
    mp=detect_columns(hs);mp.update({k:v for k,v in supplied.items() if v})
    return import_rows(cid,source,mp)

@app.get("/api/campaigns/{cid}/plan")
def campaign_plan(cid:int):
    total=one("SELECT COUNT(*) n FROM recipients WHERE campaign_id=? AND status IN ('READY','QUEUED','FAILED')",(cid,))["n"]
    ss=rows("SELECT email,daily_limit,sent_today,rate_seconds FROM senders WHERE enabled=1 ORDER BY sent_today,id")
    left=total;alloc=[]
    for s in ss:
        free=max(0,int(s["daily_limit"])-int(s["sent_today"]));take=min(free,left)
        if take:alloc.append({"sender":s["email"],"count":take,"rate_seconds":s["rate_seconds"]});left-=take
    return {"total":total,"planned_today":total-left,"remaining_after_today":left,"allocation":alloc}
@app.get("/api/recipients")
def recipients(campaign_id:int=0,status:str=""):
    q="SELECT * FROM recipients WHERE 1=1"; p=[]
    if campaign_id:q+=" AND campaign_id=?";p.append(campaign_id)
    if status:q+=" AND status=?";p.append(status)
    return rows(q+" ORDER BY id DESC LIMIT 2000",p)
@app.put("/api/recipients/{rid}")
def edit_recipient(rid:int,x:RecipientEdit):
    with con() as c:c.execute("UPDATE recipients SET email=?,company=?,name=?,subject=?,message=? WHERE id=?",(x.email,x.company,x.name,x.subject,x.message,rid))
    return {"ok":True}
@app.post("/api/campaigns/{cid}/queue")
def queue(cid:int):
    with con() as c:c.execute("UPDATE recipients SET status='QUEUED',error=NULL WHERE campaign_id=? AND status IN ('READY','FAILED')",(cid,));c.execute("UPDATE campaigns SET status='RUNNING' WHERE id=?",(cid,))
    pause_event.set(); return {"ok":True}
@app.post("/api/pause")
def pause(): pause_event.clear(); return {"ok":True}
@app.post("/api/resume")
def resume(): pause_event.set(); return {"ok":True}
@app.get("/api/logs")
def logs(): return rows("SELECT * FROM logs ORDER BY id DESC LIMIT 1000")
@app.get("/api/suppression")
def suppression(): return rows("SELECT * FROM suppression ORDER BY created_at DESC")
@app.post("/api/suppression")
def suppress(x:SuppIn):
    with con() as c:c.execute("INSERT OR REPLACE INTO suppression(email,reason,created_at) VALUES(?,?,?)",(x.email.lower(),x.reason,datetime.now().isoformat(timespec="seconds")));c.execute("UPDATE recipients SET status='UNSUBSCRIBED' WHERE lower(email)=lower(?) AND status!='SENT'",(x.email,))
    return {"ok":True}
@app.get("/api/export/{cid}")
def export(cid:int):
    rr=rows("SELECT email,company,name,subject,message,status,error FROM recipients WHERE campaign_id=?",(cid,)); out=io.StringIO(); w=csv.DictWriter(out,fieldnames=["email","company","name","subject","message","status","error"]);w.writeheader();w.writerows(rr)
    return StreamingResponse(iter([out.getvalue()]),media_type="text/csv",headers={"Content-Disposition":f"attachment; filename=campaign-{cid}.csv"})
def worker():
    while not worker_stop.is_set():
        if not pause_event.is_set(): time.sleep(.5); continue
        r=one("SELECT * FROM recipients WHERE status='QUEUED' ORDER BY id LIMIT 1")
        if not r: time.sleep(1); continue
        s=one("SELECT * FROM senders WHERE enabled=1 AND sent_today<daily_limit ORDER BY sent_today ASC,id ASC LIMIT 1")
        if not s: time.sleep(2); continue
        with con() as c:c.execute("UPDATE recipients SET status='SENDING',sender_id=?,attempts=attempts+1 WHERE id=?",(s["id"],r["id"]))
        try:
            z=smtp_connect(s); msg=EmailMessage(); msg["From"]=f'{s["name"]} <{s["email"]}>' if s["name"] else s["email"];msg["To"]=r["email"];msg["Subject"]=r["subject"];msg.set_content(r["message"]);z.send_message(msg);z.quit()
            with con() as c:c.execute("UPDATE recipients SET status='SENT',error=NULL WHERE id=?",(r["id"],));c.execute("UPDATE senders SET sent_today=sent_today+1,last_error=NULL WHERE id=?",(s["id"],))
            log(r["campaign_id"],r["id"],s["email"],r["email"],r["subject"],"SENT")
        except Exception as e:
            err=str(e)[:500]; attempts=r["attempts"]+1; status="FAILED" if attempts>=3 else "QUEUED"
            with con() as c:c.execute("UPDATE recipients SET status=?,error=? WHERE id=?",(status,err,r["id"]));c.execute("UPDATE senders SET last_error=? WHERE id=?",(err,s["id"]))
            log(r["campaign_id"],r["id"],s["email"],r["email"],r["subject"],status,err); time.sleep(min(30,2**attempts))
        time.sleep(max(.2,float(s["rate_seconds"])))
@app.on_event("startup")
def startup():
    init(); threading.Thread(target=worker,daemon=True).start()
    if getattr(sys,"frozen",False): threading.Thread(target=lambda:(time.sleep(1.2),webbrowser.open("http://127.0.0.1:8765")),daemon=True).start()
@app.on_event("shutdown")
def shutdown(): worker_stop.set()
if os.path.isdir(UI):
    assets=os.path.join(UI,"assets")
    if os.path.isdir(assets): app.mount("/assets",StaticFiles(directory=assets),name="assets")
    @app.get("/{path:path}")
    def ui(path:str):
        f=os.path.join(UI,path)
        return FileResponse(f if path and os.path.isfile(f) else os.path.join(UI,"index.html"))
if __name__=="__main__":
    import uvicorn; uvicorn.run(app,host="127.0.0.1",port=8765)
