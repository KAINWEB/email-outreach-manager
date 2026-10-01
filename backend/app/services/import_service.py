from __future__ import annotations
import csv
import io
import re
from pathlib import Path
from openpyxl import load_workbook
from email_validator import validate_email, EmailNotValidError
from sqlalchemy import delete, select
from sqlalchemy.orm import Session
from ..models import Campaign, Recipient, SuppressionEntry

FIELD_ALIASES = {
    "email": {"email", "e-mail", "mail", "почта", "почта получателя", "email получателя"},
    "company": {"company", "компания", "организация"},
    "name": {"name", "имя", "контакт", "контактное лицо"},
    "subject": {"subject", "тема", "тема письма"},
    "message": {"message", "text", "body", "сообщение", "текст", "текст получателю", "текст письма"},
}


def _normalize_header(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip().lower())


def inspect_file(path: str) -> dict:
    headers, rows = read_tabular(path, limit=5)
    suggested = {}
    for h in headers:
        n = _normalize_header(h)
        for field, aliases in FIELD_ALIASES.items():
            if n in aliases and field not in suggested:
                suggested[field] = h
    return {"headers": headers, "sample_rows": rows, "suggested_mapping": suggested}


def read_tabular(path: str, limit: int | None = None) -> tuple[list[str], list[dict[str, str]]]:
    p = Path(path)
    suffix = p.suffix.lower()
    if suffix == ".xlsx":
        wb = load_workbook(p, read_only=True, data_only=True)
        ws = wb.active
        values = ws.iter_rows(values_only=True)
        try:
            first = next(values)
        except StopIteration:
            return [], []
        headers = [str(v or "").strip() for v in first]
        rows = []
        for i, row in enumerate(values):
            item = {headers[idx]: ("" if val is None else str(val)) for idx, val in enumerate(row) if idx < len(headers)}
            rows.append(item)
            if limit is not None and i + 1 >= limit:
                break
        return headers, rows
    if suffix == ".csv":
        raw = p.read_bytes()
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw.decode("cp1251")
        sample = text[:8192]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        reader = csv.DictReader(io.StringIO(text), dialect=dialect)
        headers = reader.fieldnames or []
        rows = []
        for i, row in enumerate(reader):
            rows.append({k: (v or "") for k, v in row.items()})
            if limit is not None and i + 1 >= limit:
                break
        return headers, rows
    raise ValueError("Only .xlsx and .csv are supported")


def import_campaign(db: Session, campaign: Campaign, mapping: dict[str, str]) -> dict:
    if not campaign.source_path:
        raise ValueError("No uploaded file for this campaign")
    if "email" not in mapping:
        raise ValueError("Mapping must include the required 'email' field")
    headers, rows = read_tabular(campaign.source_path)
    missing = [col for col in mapping.values() if col and col not in headers]
    if missing:
        raise ValueError(f"Unknown source columns: {', '.join(missing)}")

    db.execute(delete(Recipient).where(Recipient.campaign_id == campaign.id))
    suppression = {x.lower() for x in db.scalars(select(SuppressionEntry.email)).all()}
    seen: set[str] = set()
    counts = {"total": 0, "ready": 0, "errors": 0, "duplicates": 0, "excluded": 0}

    for row_number, row in enumerate(rows, start=2):
        def val(field: str) -> str:
            source = mapping.get(field)
            return str(row.get(source, "") if source else "").strip()

        raw_email = val("email").lower()
        company, name, subject, message = val("company"), val("name"), val("subject"), val("message")
        status = "READY"
        errors = []
        validation_problem = False
        is_duplicate = False
        normalized_email = raw_email
        try:
            normalized_email = validate_email(raw_email, check_deliverability=False).normalized.lower()
        except EmailNotValidError as exc:
            errors.append(f"Invalid email: {exc}")
            validation_problem = True
        if normalized_email in seen:
            is_duplicate = True
            status = "SKIPPED"
            errors.append("Duplicate within campaign")
            counts["duplicates"] += 1
        seen.add(normalized_email)
        if normalized_email in suppression:
            status = "UNSUBSCRIBED"
            errors.append("Address is in suppression list")
            counts["excluded"] += 1
        if not subject:
            errors.append("Missing subject")
            validation_problem = True
        if not message:
            errors.append("Missing message")
            validation_problem = True
        if validation_problem and status == "READY":
            status = "SKIPPED"
        if validation_problem:
            counts["errors"] += 1
        if status == "READY":
            counts["ready"] += 1
        counts["total"] += 1
        db.add(Recipient(
            campaign_id=campaign.id,
            row_number=row_number,
            email=normalized_email,
            company=company,
            name=name,
            subject=subject,
            message=message,
            status=status,
            validation_error="; ".join(errors) if errors else None,
            is_duplicate=is_duplicate,
        ))
    campaign.status = "REVIEW"
    db.commit()
    return counts


def revalidate_recipient(db: Session, recipient: Recipient) -> Recipient:
    errors = []
    try:
        recipient.email = validate_email(recipient.email, check_deliverability=False).normalized.lower()
    except EmailNotValidError as exc:
        errors.append(f"Invalid email: {exc}")
    if not recipient.subject.strip():
        errors.append("Missing subject")
    if not recipient.message.strip():
        errors.append("Missing message")
    if db.scalar(select(SuppressionEntry).where(SuppressionEntry.email == recipient.email)):
        recipient.status = "UNSUBSCRIBED"
        errors.append("Address is in suppression list")
    elif errors:
        recipient.status = "SKIPPED"
    else:
        duplicate = db.scalar(select(Recipient).where(
            Recipient.campaign_id == recipient.campaign_id,
            Recipient.email == recipient.email,
            Recipient.id != recipient.id,
            Recipient.status.in_(["READY", "QUEUED", "SENDING", "SENT"]),
        ))
        recipient.is_duplicate = duplicate is not None
        if duplicate:
            recipient.status = "SKIPPED"
            errors.append("Duplicate within campaign")
        else:
            recipient.status = "READY"
    recipient.validation_error = "; ".join(errors) if errors else None
    db.commit()
    db.refresh(recipient)
    return recipient
